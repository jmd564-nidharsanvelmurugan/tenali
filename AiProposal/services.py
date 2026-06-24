import io
import re
import time
import uuid
from fastapi.responses import StreamingResponse
import requests
from sqlalchemy.orm import Session
from sqlalchemy import select
from uuid import UUID
from db.models import ProposalQuestions, Message, ProposalMetadata, Conversation, Workspace, User
from .workspace_integration import (
    get_ai_proposal_workspace, 
    create_ai_proposal_conversation,
    ensure_ai_proposal_workspace_exists,
    is_ai_proposal_conversation
)
from .schemas import ProposalMetadataSchema, ScoredProposals, EditAnswersRequest
from .utils.input_parser import *
from sqlalchemy import asc, text
from .utils.get_n_matching_proposals import *
from .utils.prompts import *
import os
import asyncio
from fastapi import FastAPI, HTTPException, UploadFile
from openai import AzureOpenAI
import httpx
from azure.storage.blob import BlobServiceClient
from azure.storage.blob.aio import BlobServiceClient as AsyncBlobServiceClient
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.document import Document as _Document
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from markdown2 import markdown
from bs4 import BeautifulSoup
import json
from .utils.prompts import AI_PROPOSAL_SYSTEM_PROMPT, DOCS_PROMPT
from fastapi import status
from fastapi.responses import JSONResponse
from datetime import datetime, timezone, timedelta
from .Langgraph.main import generate_proposal_langgraph


# Azure Search configuration
AZURE_SEARCH_QUERY_TYPE = os.getenv("AZURE_SEARCH_QUERY_TYPE", "simple")
AZURE_SEARCH_USE_SEMANTIC_SEARCH = os.getenv("AZURE_SEARCH_USE_SEMANTIC_SEARCH", "false")
AZURE_SEARCH_SEMANTIC_SEARCH_CONFIG = os.getenv("AZURE_SEARCH_SEMANTIC_SEARCH_CONFIG", "")
AZURE_OPENAI_TEMPERATURE = os.getenv("AZURE_OPENAI_TEMPERATURE", "0.7")
AZURE_OPENAI_MAX_TOKENS = os.getenv("AZURE_OPENAI_MAX_TOKENS", "1000")
AZURE_OPENAI_TOP_P = os.getenv("AZURE_OPENAI_TOP_P", "1.0")
AZURE_OPENAI_STOP_SEQUENCE = os.getenv("AZURE_OPENAI_STOP_SEQUENCE", "")
AZURE_SEARCH_SERVICE = os.getenv("AZURE_SEARCH_SERVICE")
AZURE_SEARCH_KEY = os.getenv("AZURE_SEARCH_KEY")
AZURE_SEARCH_CONTENT_COLUMNS = os.getenv("AZURE_SEARCH_CONTENT_COLUMNS", "")
AZURE_SEARCH_TITLE_COLUMN = os.getenv("AZURE_SEARCH_TITLE_COLUMN", "")
AZURE_SEARCH_URL_COLUMN = os.getenv("AZURE_SEARCH_URL_COLUMN", "")
AZURE_SEARCH_FILENAME_COLUMN = os.getenv("AZURE_SEARCH_FILENAME_COLUMN", "")
AZURE_SEARCH_VECTOR_COLUMNS = os.getenv("AZURE_SEARCH_VECTOR_COLUMNS", "")
AZURE_SEARCH_TOP_K = int(os.getenv("AZURE_SEARCH_TOP_K", 5))
AZURE_OPENAI_SYSTEM_MESSAGE = os.getenv("AZURE_OPENAI_SYSTEM_MESSAGE", "You are a helpful assistant.")
AZURE_OPENAI_EMBEDDING_ENDPOINT = os.getenv("AZURE_OPENAI_EMBEDDING_ENDPOINT")
AZURE_OPENAI_EMBEDDING_KEY = os.getenv("AZURE_OPENAI_EMBEDDING_KEY")
AZURE_SEARCH_STRICTNESS = os.getenv("AZURE_SEARCH_STRICTNESS", "0")
AZURE_SEARCH_SERVICE_INDEXER_GENERATED_PROPOSALS_NAME = os.getenv("AZURE_SEARCH_SERVICE_INDEXER_GENERATED_PROPOSALS_NAME", "generated-proposals-indexer")

AZURE_SEARCH_AI_PROPOSAL_WORKSPACE_INDEXER = os.getenv("AZURE_SEARCH_AI_PROPOSAL_WORKSPACE_INDEXER", "ai-proposal-workspace-indexer")
AI_PROPOSAL_WORKSPACE_CONTAINER_NAME = os.getenv("AI_PROPOSAL_WORKSPACE_CONTAINER_NAME", "workspaces")
AZURE_OPENAI_KEY = os.getenv("AZURE_OPENAI_KEY")
AZURE_OPENAI_ENDPOINT = os.getenv("AZURE_OPENAI_ENDPOINT")
AZURE_OPENAI_DEPLOYMENT = os.getenv("AZURE_OPENAI_DEPLOYMENT_MODEL_NAME_1")
AZURE_OPENAI_PREVIEW_API_VERSION = os.getenv("AZURE_OPENAI_PREVIEW_API_VERSION", "2024-05-01-preview")
PROPOSAL_TEMPLATE_CONTAINER_NAME = os.getenv("PROPOSAL_TEMPLATE_CONTAINER_NAME", "workspaces")
GENERATED_PROPOSALS_CONTAINER_NAME = os.getenv("GENERATED_PROPOSALS_CONTAINER_NAME", "workspaces")
AI_PROPOSAL_SALES_QAS_CONTAINER_NAME = os.getenv("AI_PROPOSAL_SALES_QAS_CONTAINER_NAME", "ai-proposal-sales-qas")

blob_service_client = BlobServiceClient.from_connection_string(
    os.environ.get("AZURE_BLOB_STORAGE_CONNECTION_STRING")
)

headers = {
    "api-key": AZURE_OPENAI_KEY,
    "Content-Type": "application/json",
}

httpx_client = httpx.Client(
    base_url=AZURE_OPENAI_ENDPOINT,
    headers={"api-key": AZURE_OPENAI_KEY},
    timeout=60.0,
)

azure_client = AzureOpenAI(
    api_key=AZURE_OPENAI_KEY,
    api_version="2024-05-01-preview",
    azure_endpoint=AZURE_OPENAI_ENDPOINT,
    http_client=httpx_client,
)


async def update_message_service(request, db):
    """Update an existing message's content in the database"""
    message = db.query(Message).filter(
        Message.conversation_id == request.conversation_id,
        Message.id == request.msg_id
    ).first()
    if not message:
        from fastapi import HTTPException
        raise HTTPException(status_code=404, detail="Message not found")
    message.content = request.content
    db.commit()
    return {"message": f"Message {request.msg_id} updated successfully"}


async def fetch_chat_completion(prompt: str, system_prompt: str, docs_prompt: str, conversation_id: str, data_source=None) -> dict:
    """Send a chat completion request using Azure OpenAI client"""
    messages = [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": prompt}
    ]
    try:
        response = azure_client.chat.completions.create(
            model=AZURE_OPENAI_DEPLOYMENT,
            messages=messages,
            temperature=float(AZURE_OPENAI_TEMPERATURE),
            max_tokens=10000,
            top_p=float(AZURE_OPENAI_TOP_P),
        )
        return {
            "choices": [{
                "message": {
                    "content": response.choices[0].message.content
                }
            }]
        }
    except Exception as e:
        print(f"Azure OpenAI error: {e}")
        raise HTTPException(status_code=500, detail=f"AI generation failed: {str(e)}")


def iter_block_items(parent):
    """Iterate through paragraph and table blocks in a Word document"""
    if isinstance(parent, _Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise ValueError("Unsupported parent type.")
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def get_all_text_from_doc(doc):
    """Extract all text content from a Word document"""
    parts = []
    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            if block.text.strip():
                parts.append(block.text.strip())
        elif isinstance(block, Table):
            for row in block.rows:
                row_text = [cell.text.strip() for cell in row.cells]
                if any(row_text):
                    parts.append("\t".join(row_text))
    return "\n".join(parts)


def get_proposal_content(proposals: list[ScoredProposals]) -> list[str]:
    """Retrieve content from proposal template files stored in blob storage"""
    combined_content = []
    for p in proposals:
        try:
            container_name = PROPOSAL_TEMPLATE_CONTAINER_NAME or "workspaces"
            blob_client = blob_service_client.get_blob_client(
                container=container_name,
                blob='templates/' + p.proposal.name
            )
            blob_bytes = blob_client.download_blob().readall()
            doc = Document(io.BytesIO(blob_bytes))
            text = get_all_text_from_doc(doc)
            combined_content.append(text)
        except Exception as e:
            print(f"Warning: Could not load proposal template {p.proposal.name}: {e}")
            combined_content.append(f"Template content for {p.proposal.name} not available")
    return combined_content


async def read_sales_call_questions_docx():
    """Read and return the sales call questions DOCX template from blob storage"""
    blob_name = "templates/sales_call_questions_v1.docx"
    blob_client = blob_service_client.get_blob_client(
        container=PROPOSAL_TEMPLATE_CONTAINER_NAME,
        blob=blob_name
    )
    blob_bytes = blob_client.download_blob().readall()
    doc_stream = io.BytesIO(blob_bytes)
    doc_stream.seek(0)
    return StreamingResponse(
        doc_stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=sales_call_questions.docx"}
    )


async def upload_sales_call_questions_docx(user_id: UUID, conversation_id: UUID, file: UploadFile):
    """Upload a sales call questions DOCX file to blob storage"""
    try:
        blob_service_client_async = AsyncBlobServiceClient.from_connection_string(
            os.environ.get("AZURE_BLOB_STORAGE_CONNECTION_STRING")
        )
        container_client = blob_service_client_async.get_container_client(AI_PROPOSAL_SALES_QAS_CONTAINER_NAME)
        blob_name = f"{user_id}----{conversation_id}.docx"
        await container_client.upload_blob(blob_name, await file.read(), overwrite=True)
        return JSONResponse(
            status_code=status.HTTP_201_CREATED,
            content={"message": blob_name}
        )
    except Exception as e:
        return JSONResponse(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            content={"message": f"File upload failed: {str(e)}"}
        )


async def read_uploaded_sales_call_questions_docx(user_id: UUID, conversation_id: UUID):
    """Read the uploaded sales call questions DOCX from blob storage. Returns empty string if not present."""
    blob_name = f"{user_id}----{conversation_id}.docx"
    try:
        blob_service_client_async = AsyncBlobServiceClient.from_connection_string(
            os.environ.get("AZURE_BLOB_STORAGE_CONNECTION_STRING")
        )
        blob_client = blob_service_client_async.get_blob_client(
            container=AI_PROPOSAL_SALES_QAS_CONTAINER_NAME,
            blob=blob_name
        )
        blob_bytes = await (await blob_client.download_blob()).readall()
        doc = Document(io.BytesIO(blob_bytes))
        text = get_all_text_from_doc(doc)
        return text
    except Exception:
        return ""


async def generate_proposal(conversation_id: UUID, db: Session, uid: UUID, user_prompt: str = None):
    """
    Generate a proposal based on uploaded sales call DOCX.
    
    Reads the uploaded questionnaire DOCX from blob storage and passes it
    to the langgraph agent. The conversation does not need any pre-existing
    messages — workspace_id is sourced directly from the Conversation record.
    """
    # Get current user
    current_user = db.query(User).filter_by(id=uid).first()
    if not current_user:
        raise ValueError(f"User {uid} not found")

    # Get conversation (needed for workspace_id)
    conversation = db.query(Conversation).filter_by(id=conversation_id).first()
    if not conversation:
        raise ValueError(f"Conversation {conversation_id} not found")

    # Read the uploaded sales call DOCX from blob storage
    client_context_qas = await read_uploaded_sales_call_questions_docx(uid, conversation_id)

    if client_context_qas:
        questionnaire_for_agent = client_context_qas
    else:
        questionnaire_for_agent = ""

    # Call generate_proposal_langgraph() with the resolved questionnaire + user prompt
    ggg_output = generate_proposal_langgraph(
        questionnaire=questionnaire_for_agent,
        user_prompt=user_prompt or "",
    )

    # Extract data from ggg_output
    proposal_text = ggg_output.get("proposal_text", "")
    sections = ggg_output.get("sections", [])

    citations_data = ggg_output.get("citations", {})
    if isinstance(citations_data, dict) and "citations" in citations_data:
        citations_list = citations_data["citations"]
    elif isinstance(citations_data, list):
        citations_list = citations_data
    else:
        citations_list = []

    if not isinstance(citations_list, list):
        citations_list = []

    # Store generated proposal to DB
    now = datetime.now(timezone.utc)
    delta = timedelta(milliseconds=1)

    new_message = Message(
        user_id=uid,
        conversation_id=conversation_id,
        workspace_id=conversation.workspace_id,
        content=proposal_text,
        role="assistant",
        model_type="",
        chat_type="hybrid",
        input_tokens=0,
        output_tokens=0,
        created_at=now,
    )

    citations_for_db = {"citations": citations_list}
    reference_proposals = Message(
        user_id=uid,
        conversation_id=conversation_id,
        workspace_id=conversation.workspace_id,
        content=json.dumps(citations_for_db),
        role="tool",
        model_type="",
        chat_type="hybrid",
        input_tokens=0,
        output_tokens=0,
        created_at=now + delta,
    )

    db.add(new_message)
    db.add(reference_proposals)
    db.commit()

    # Upload proposal to blob storage
    try:
        proposal_docx_upload(proposal_text, current_user.email, conversation_id)
    except Exception as e:
        print(f"Warning: Could not upload to blob storage: {e}")

    # Format sections for frontend
    formatted_sections = []

    if sections and len(sections) > 0:
        for section in sections:
            if isinstance(section, dict):
                formatted_sections.append({
                    "prompt": section.get("prompt", "Section"),
                    "response": section.get("response", ""),
                    "error": section.get("error")
                })
            else:
                formatted_sections.append({
                    "prompt": "Section",
                    "response": str(section),
                    "error": None
                })
    elif proposal_text:
        formatted_sections = parse_sections_from_text(proposal_text)
    else:
        formatted_sections = [{
            "prompt": "Proposal",
            "response": "No content generated",
            "error": None
        }]

    return {
        "msg_id": new_message.id,
        "proposal": formatted_sections,
        "citations": {
            "citations": [
                {
                    "filepath": c.get("filepath", ""),
                    "url": c.get("url", "")
                }
                for c in citations_list if isinstance(c, dict)
            ]
        }
    }


def edit_answers(request: EditAnswersRequest, db: Session):
    """Edit answers in the conversation based on provided question IDs and new content"""
    try:
        conv_messages = (
            db.query(Message)
            .filter(Message.conversation_id == request.conversation_id)
            .order_by(asc(Message.created_at))
            .all()
        )
        for item in request.to_edit:
            qid_index = int(item.qid) * 2 - 1
            if qid_index < len(conv_messages):
                conv_messages[qid_index].content = item.content
            else:
                raise IndexError(f"Invalid qid: {item.qid}, out of range for conversation messages.")
        db.commit()
        return {"message": "success"}
    except Exception as e:
        db.rollback()
        return {"message": str(e)}


async def follow_up_proposal_service(conversation_id: UUID, recent_message: str, new_message: str, db: Session, current_user: User):
    """Generate a follow-up proposal based on new instructions"""
    system_prompt = """Update the entire proposal that is provided based on the new instruction provided.
Give the full updated proposal. Don't add anything other than updated proposal in the response
STRICTLY generate the proposal content in a WELL-STRUCTURED MARKDOWN FORMAT, use h1 heading style for section headings, h2, h3 for appropriate sub-headings, ul for bullet points and ol for numbered lists.
"""
    prompt = (
        f"Proposal:\n{recent_message}\n"
        f"New instruction:\n{new_message}"
    )
    docs_prompt = """
Use the documents as a reference to generate the updated proposal.
Adhere to the user and system instructions as well.
"""
    try:
        llm_response = await fetch_chat_completion(
            prompt, system_prompt, docs_prompt, str(conversation_id)
        )

        assistant_content = llm_response["choices"][0]["message"]["content"]
        tool_content = json.dumps({"citations": []})

        last_message = (
            db.query(Message)
            .filter(Message.conversation_id == conversation_id)
            .order_by(Message.created_at.desc())
            .first()
        )
        if not last_message:
            return {"message": "No message found to update"}

        now = datetime.now(timezone.utc)
        delta = timedelta(milliseconds=1)

        user_msg = Message(
            user_id=last_message.user_id,
            conversation_id=conversation_id,
            workspace_id=last_message.workspace_id,
            content=new_message,
            role="user",
            model_type="",
            chat_type="hybrid",
            input_tokens=0,
            output_tokens=0,
            created_at=now,
        )
        assistant_msg = Message(
            user_id=last_message.user_id,
            conversation_id=conversation_id,
            workspace_id=last_message.workspace_id,
            content=assistant_content,
            role="assistant",
            model_type="",
            chat_type="hybrid",
            input_tokens=0,
            output_tokens=0,
            created_at=now + delta,
        )
        tool_msg = Message(
            user_id=last_message.user_id,
            conversation_id=conversation_id,
            workspace_id=last_message.workspace_id,
            content=tool_content,
            role="tool",
            model_type="",
            chat_type="hybrid",
            input_tokens=0,
            output_tokens=0,
            created_at=now + delta + delta,
        )
        db.add(user_msg)
        db.add(assistant_msg)
        db.add(tool_msg)
        db.commit()

        return {
            "msg_id": assistant_msg.id,
            "assistant": assistant_content,
            "tool": tool_content,
            "replace_existing": True,
            "target_message_id": str(user_msg.id)
        }
    except Exception as e:
        db.rollback()
        return {"message": str(e)}


def add_content_to_doc(doc, content):
    """Add HTML or Markdown content to a Word document with proper formatting"""
    if "<" not in content and ">" not in content:
        html_content = markdown(content, extras=["fenced-code-blocks"])
    else:
        html_content = content

    soup = BeautifulSoup(html_content, "html.parser")

    for element in soup:
        if element.name == "h1":
            doc.add_paragraph(element.get_text(), style="Heading 1")
        elif element.name == "h2":
            doc.add_paragraph(element.get_text(), style="Heading 2")
        elif element.name == "h3":
            doc.add_paragraph(element.get_text(), style="Heading 3")
        elif element.name == "ul":
            for li in element.find_all("li"):
                p = doc.add_paragraph()
                p.add_run("• ").bold = False
                p.add_run(li.get_text())
        elif element.name == "ol":
            counter = 1
            for li in element.find_all("li"):
                p = doc.add_paragraph()
                p.add_run(f"{counter}. ").bold = False
                p.add_run(li.get_text())
                counter += 1
        elif element.name == "p":
            doc.add_paragraph(element.get_text())
        elif element.name == "strong":
            p = doc.add_paragraph()
            p.add_run(element.get_text()).bold = True
        elif element.name == "em":
            p = doc.add_paragraph()
            p.add_run(element.get_text()).italic = True
        elif element.name == "code":
            p = doc.add_paragraph(style="Quote")
            p.add_run(element.get_text()).font.name = "Courier New"
        elif element.name == "br":
            doc.add_paragraph("")

    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT


def proposal_docx(conversation_id: UUID, db: Session):
    """
    Generate a formatted DOCX proposal document from the conversation.
    Creates a professional proposal with cover page, headers, footers, and styled content.
    """
    try:
        import os
        import re
        import json
        import io
        import base64
        from datetime import datetime
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor, Cm, Emu
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
        from docx.enum.section import WD_SECTION
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        from typing import Optional, Dict, Any, List
        from uuid import UUID
        from sqlalchemy import asc
        from PIL import Image
        from io import BytesIO

        # Import Azure agent and Gantt tools
        from .utils.azure_agent import get_agent_response
        from .utils.gantt import create_hierarchical_gantt

        _NAVY    = RGBColor(0x17, 0x38, 0x45)
        _PINK    = RGBColor(0xFF, 0x61, 0x96)
        _DARK    = RGBColor(0x1D, 0x1C, 0x1C)
        _GRAY    = RGBColor(0x80, 0x80, 0x80)
        _DEEP    = RGBColor(0x19, 0x10, 0x5B)
        _WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
        _BORDER  = RGBColor(0xDD, 0xDD, 0xDD)
        _RED     = RGBColor(0xCC, 0x00, 0x00)
        _AMBER   = RGBColor(0xCC, 0x77, 0x00)
        _GREEN   = RGBColor(0x2E, 0x7D, 0x32)
        _FONT = "Arial"
        _ZEBRA = "F9F9F9"

        _PAGE_W_CM = 21.0
        _PAGE_H_CM = 29.7

        # Asset directory (local files - kept for reading)
        _ASSETS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "jman")
        _ASSET_FILES = {
            "logo_white":  "jman_logo_white.png",
            "logo_navy":   "jman_logo_navy.png",
            "corner_mark": "jman_corner_mark.png",
            "cover_bg":    "cover_background.jpg",
        }

        def _asset(name: str):
            path = os.path.join(_ASSETS_DIR, _ASSET_FILES[name])
            return path if os.path.isfile(path) else None

        def _run(para, text, size_pt, bold=False, italic=False, color=None, font=_FONT):
            run = para.add_run(text)
            run.font.name = font
            run.font.size = Pt(size_pt)
            run.font.bold = bold
            run.font.italic = italic
            if color:
                run.font.color.rgb = color
            return run

        def _set_cell_shading(cell, fill_hex: str):
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            for old in tcPr.findall(qn("w:shd")):
                tcPr.remove(old)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), fill_hex)
            tcPr.append(shd)

        def _set_cell_borders(cell, color_hex="DDDDDD", sz=4):
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            for old in tcPr.findall(qn("w:tcBorders")):
                tcPr.remove(old)
            tcBorders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                el = OxmlElement(f"w:{side}")
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), str(sz))
                el.set(qn("w:space"), "0")
                el.set(qn("w:color"), color_hex)
                tcBorders.append(el)
            tcPr.append(tcBorders)

        def _set_cell_margins(cell, top=80, bottom=80, left=110, right=110):
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            mar = OxmlElement("w:tcMar")
            for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
                el = OxmlElement(f"w:{side}")
                el.set(qn("w:w"), str(val))
                el.set(qn("w:type"), "dxa")
                mar.append(el)
            tcPr.append(mar)

        def _set_table_width(tbl, width_twips: int):
            tbl_el = tbl._tbl
            tblPr = tbl_el.find(qn("w:tblPr"))
            if tblPr is None:
                tblPr = OxmlElement("w:tblPr")
                tbl_el.insert(0, tblPr)
            for old in tblPr.findall(qn("w:tblW")):
                tblPr.remove(old)
            tblW = OxmlElement("w:tblW")
            tblW.set(qn("w:w"), str(width_twips))
            tblW.set(qn("w:type"), "dxa")
            tblStyle = tblPr.find(qn("w:tblStyle"))
            if tblStyle is not None:
                tblStyle.addnext(tblW)
            else:
                tblPr.insert(0, tblW)

        def _pink_divider(doc):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "8")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "FF6196")
            pBdr.append(bottom)
            pPr.append(pBdr)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(10)
            return p

        def _spacer(doc, pt=8):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(pt)
            return p

        def _add_floating_picture(paragraph, image_path, width, height, behind_doc=True, name="Picture"):
            run = paragraph.add_run()
            run.add_picture(image_path, width=width, height=height)
            drawing = run._r.find(qn("w:drawing"))
            inline = drawing.find(qn("wp:inline"))
            extent = inline.find(qn("wp:extent"))
            docPr = inline.find(qn("wp:docPr"))
            graphic = inline.find(qn("a:graphic"))
            inline.remove(extent)
            inline.remove(docPr)
            inline.remove(graphic)
            docPr.set("name", name)

            anchor = OxmlElement("wp:anchor")
            anchor.set("distT", "0")
            anchor.set("distB", "0")
            anchor.set("distL", "0")
            anchor.set("distR", "0")
            anchor.set("simplePos", "0")
            anchor.set("relativeHeight", "251658240")
            anchor.set("behindDoc", "1" if behind_doc else "0")
            anchor.set("locked", "0")
            anchor.set("layoutInCell", "1")
            anchor.set("allowOverlap", "1")

            simplePos = OxmlElement("wp:simplePos")
            simplePos.set("x", "0")
            simplePos.set("y", "0")

            posH = OxmlElement("wp:positionH")
            posH.set("relativeFrom", "page")
            offH = OxmlElement("wp:posOffset")
            offH.text = "0"
            posH.append(offH)

            posV = OxmlElement("wp:positionV")
            posV.set("relativeFrom", "page")
            offV = OxmlElement("wp:posOffset")
            offV.text = "0"
            posV.append(offV)

            effectExtent = OxmlElement("wp:effectExtent")
            for side in ("l", "t", "r", "b"):
                effectExtent.set(side, "0")

            wrapNone = OxmlElement("wp:wrapNone")

            anchor.append(simplePos)
            anchor.append(posH)
            anchor.append(posV)
            anchor.append(extent)
            anchor.append(effectExtent)
            anchor.append(wrapNone)
            anchor.append(docPr)
            anchor.append(graphic)

            drawing.remove(inline)
            drawing.append(anchor)
            return run

        def _add_page_number_field(paragraph, size_pt=8, color=_GRAY):
            run = paragraph.add_run()
            run.font.size = Pt(size_pt)
            run.font.color.rgb = color
            fld = OxmlElement('w:fld')
            instr = OxmlElement('w:instrText')
            instr.text = 'PAGE'
            fld.append(instr)
            run._r.append(fld)

        # =====================================================
        # Helper: Generate Gantt chart data from Approach content
        # =====================================================
        def _generate_gantt_data(content: str) -> List[Dict[str, Any]]:
            """Extract phase and task data from Approach content to generate Gantt chart data."""
            system_prompt = """
            You are a data extraction specialist. Extract phase and task information from the given Approach section.
            
            Parse the Approach content and generate a JSON array of phases with their tasks.
            
            Each phase should have:
            - Name: Phase name with number
            - Start_Day: Starting day number (1-based)
            - Finish_Day: Finishing day number
            - Type: "Phase"
            - Progress: 0-100 (estimate based on task completion)
            - isMain: True
            - tasks: Array of task objects
            
            Each task should have:
            - Name: Task name
            - Start_Day: Starting day number
            - Finish_Day: Finishing day number
            - Type: "Task"
            - Progress: 0-100 (estimate based on phase progress)
            
            Use the timeline information from the Approach section (Week Y to Week Z) to calculate day numbers.
            Assume each week = 5 working days.
            
            Return ONLY the JSON array, no other text.
            """
            
            prompt = f"""
            Extract phase and task data from this Approach section:
            
            {content}
            
            Generate a JSON array of phases with tasks in the exact format:
            [
                {{
                    "Name": "Phase 1 - [Name]",
                    "Start_Day": 1,
                    "Finish_Day": 21,
                    "Type": "Phase",
                    "Progress": 100,
                    "isMain": True,
                    "tasks": [
                        {{
                            "Name": "Task 1 - [Name]",
                            "Start_Day": 1,
                            "Finish_Day": 7,
                            "Type": "Task",
                            "Progress": 100
                        }}
                    ]
                }}
            ]
            
            Calculate day numbers from the timeline (Week X to Week Y).
            Assume 5 working days per week.
            """
            
            response = get_agent_response(prompt, system_prompt)
            
            try:
                json_str = response.strip()
                if "```json" in json_str:
                    start = json_str.find("```json") + 7
                    end = json_str.rfind("```")
                    json_str = json_str[start:end].strip()
                elif "```" in json_str:
                    start = json_str.find("```") + 3
                    end = json_str.rfind("```")
                    json_str = json_str[start:end].strip()
                
                data = json.loads(json_str)
                print(f"✅ Successfully extracted {len(data)} phases from Approach content")
                return data
            except json.JSONDecodeError as e:
                print(f"❌ Failed to parse Gantt data: {e}")
                # Return default structure
                return [
                    {
                        "Name": "Phase 1 - Planning",
                        "Start_Day": 1,
                        "Finish_Day": 21,
                        "Type": "Phase",
                        "Progress": 100,
                        "isMain": True,
                        "tasks": [
                            {"Name": "Requirements", "Start_Day": 1, "Finish_Day": 7, "Type": "Task", "Progress": 100},
                            {"Name": "Analysis", "Start_Day": 4, "Finish_Day": 14, "Type": "Task", "Progress": 100},
                            {"Name": "Approval", "Start_Day": 15, "Finish_Day": 21, "Type": "Task", "Progress": 100}
                        ]
                    },
                    {
                        "Name": "Phase 2 - Development",
                        "Start_Day": 22,
                        "Finish_Day": 45,
                        "Type": "Phase",
                        "Progress": 50,
                        "isMain": True,
                        "tasks": [
                            {"Name": "Backend", "Start_Day": 22, "Finish_Day": 30, "Type": "Task", "Progress": 70},
                            {"Name": "Frontend", "Start_Day": 31, "Finish_Day": 40, "Type": "Task", "Progress": 40},
                            {"Name": "Integration", "Start_Day": 38, "Finish_Day": 45, "Type": "Task", "Progress": 20}
                        ]
                    }
                ]

        # =====================================================
        # Helper: Add base64 image to document
        # =====================================================
        def _add_base64_image_to_doc(doc, base64_string: str, width_cm: float = 15.0):
            """Add an image from a base64 string to the document."""
            if not base64_string:
                print("   ⚠️ No base64 image data provided")
                return False
            
            try:
                image_data = base64.b64decode(base64_string)
                image_stream = BytesIO(image_data)
                img = Image.open(image_stream)
                
                aspect_ratio = img.height / img.width
                height_cm = width_cm * aspect_ratio
                image_stream.seek(0)
                
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(6)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                
                run = p.add_run()
                run.add_picture(image_stream, width=Cm(width_cm), height=Cm(height_cm))
                
                caption = doc.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _run(caption, "Figure: Project Timeline - Gantt Chart", 9, italic=True, color=_GRAY)
                caption.paragraph_format.space_before = Pt(4)
                caption.paragraph_format.space_after = Pt(12)
                
                print(f"   ✅ Added Gantt chart image (width: {width_cm}cm, height: {height_cm:.2f}cm)")
                return True
                
            except Exception as e:
                print(f"   ❌ Failed to add base64 image: {e}")
                return False

        _BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
        _SECTION_DIVIDER_RE = re.compile(r"^---\s*$|^___\s*$|^\*\*\*\s*$")

        def _add_inline_runs(p, text, size_pt, color, bold=False, italic=False):
            pos = 0
            for m in _BOLD_RE.finditer(text):
                if m.start() > pos:
                    _run(p, text[pos:m.start()], size_pt, bold=bold, italic=italic, color=color)
                _run(p, m.group(1), size_pt, bold=True, italic=italic, color=color)
                pos = m.end()
            if pos < len(text):
                _run(p, text[pos:], size_pt, bold=bold, italic=italic, color=color)

        def _status_color(value: str):
            v = value.strip().lower()
            if any(k in v for k in ("blocked", "at risk", "failed", "delayed")):
                return _RED
            if any(k in v for k in ("in progress", "pending", "ongoing", "review")):
                return _AMBER
            if any(k in v for k in ("resolved", "done", "complete", "approved", "on track")):
                return _GREEN
            return _DARK

        def _render_table(doc, header_cells, body_rows):
            n_cols = len(header_cells)
            tbl = doc.add_table(rows=1, cols=n_cols)
            tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
            tbl.style = "Table Grid"

            content_width_twips = int((_PAGE_W_CM - 3.5) * 567)
            if n_cols == 2:
                widths = [int(content_width_twips * 0.7), int(content_width_twips * 0.3)]
            else:
                widths = [content_width_twips // n_cols] * n_cols
            _set_table_width(tbl, sum(widths))

            hdr_row = tbl.rows[0]
            for i, text in enumerate(header_cells):
                cell = hdr_row.cells[i]
                cell.width = Pt(widths[i] / 20)
                _set_cell_shading(cell, "173845")
                _set_cell_borders(cell, "DDDDDD", 4)
                _set_cell_margins(cell)
                p = cell.paragraphs[0]
                for r in list(p.runs):
                    r._r.getparent().remove(r._r)
                _run(p, text, 10, bold=True, color=_WHITE)

            for ridx, row_vals in enumerate(body_rows):
                row = tbl.add_row()
                fill = "FFFFFF" if ridx % 2 == 0 else _ZEBRA
                for i, val in enumerate(row_vals):
                    cell = row.cells[i]
                    cell.width = Pt(widths[i] / 20)
                    _set_cell_shading(cell, fill)
                    _set_cell_borders(cell, "DDDDDD", 4)
                    _set_cell_margins(cell)
                    p = cell.paragraphs[0]
                    is_status_col = (n_cols == 2 and i == 1)
                    color = _status_color(val) if is_status_col else _DARK
                    _run(p, val, 10, bold=is_status_col, color=color)
            return tbl

        def _apply_styled_content(doc, raw_text: str, heading_counter: list, extract_gantt: bool = False):
            """
            Apply styled content to document.
            If extract_gantt is True, this is the Approach section and we should generate Gantt chart.
            """
            lines = raw_text.splitlines()
            i = 0
            n = len(lines)

            # Collect the approach content to generate Gantt from
            approach_content = raw_text

            while i < n:
                line = lines[i]
                stripped = line.strip()

                if not stripped:
                    doc.add_paragraph().paragraph_format.space_after = Pt(4)
                    i += 1
                    continue

                if _SECTION_DIVIDER_RE.match(stripped):
                    _pink_divider(doc)
                    i += 1
                    continue

                if stripped.startswith("|") and i + 1 < n and re.match(r"^\|?[\s:|-]+\|?$", lines[i + 1].strip()):
                    header_cells = [c.strip() for c in stripped.strip("|").split("|")]
                    j = i + 2
                    body_rows = []
                    while j < n and lines[j].strip().startswith("|"):
                        body_rows.append([c.strip() for c in lines[j].strip().strip("|").split("|")])
                        j += 1
                    _render_table(doc, header_cells, body_rows)
                    _spacer(doc, pt=8)
                    i = j
                    continue

                if stripped.startswith("# "):
                    if heading_counter[0] > 0:
                        _pink_divider(doc)
                    heading_counter[0] += 1
                    text = stripped[2:].strip()
                    p = doc.add_paragraph()
                    pf = p.paragraph_format
                    pf.left_indent = Cm(1.02)
                    pf.first_line_indent = Cm(-1.02)
                    pf.space_before = Pt(14)
                    pf.space_after = Pt(8)
                    try:
                        pf.tab_stops.add_tab_stop(Cm(1.02))
                    except Exception:
                        pass
                    _run(p, f"{heading_counter[0]}\t", 12, bold=True, color=_DEEP)
                    _add_inline_runs(p, text, 12, _DEEP, bold=True)
                    i += 1
                    continue

                if stripped.startswith("## "):
                    text = stripped[3:].strip()
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(10)
                    p.paragraph_format.space_after = Pt(4)
                    _add_inline_runs(p, text, 11, _PINK, bold=True)
                    i += 1
                    continue

                if stripped.startswith("- ") or stripped.startswith("* "):
                    text = stripped[2:].strip()
                    p = doc.add_paragraph()
                    pf = p.paragraph_format
                    pf.left_indent = Cm(1.27)
                    pf.first_line_indent = Cm(-0.64)
                    pf.space_after = Pt(4)
                    pf.line_spacing = 1.15
                    try:
                        pf.tab_stops.add_tab_stop(Cm(1.27))
                    except Exception:
                        pass
                    _run(p, "•\t", 10, color=_PINK)
                    _add_inline_runs(p, text, 10, _DARK)
                    i += 1
                    continue

                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.line_spacing = 1.15
                _add_inline_runs(p, stripped, 10, _DARK)
                i += 1

            # After rendering all content, if this is the Approach section, generate and add Gantt chart
            if extract_gantt and approach_content:
                print("\n   📊 Generating Gantt chart for Approach section...")
                try:
                    # Generate Gantt data from the approach content
                    gantt_data = _generate_gantt_data(approach_content)
                    
                    if gantt_data:
                        # Create the base64 encoded Gantt chart
                        base64_gantt = create_hierarchical_gantt(gantt_data)
                        
                        if base64_gantt:
                            # Add the image to the document
                            _add_base64_image_to_doc(doc, base64_gantt, width_cm=15.0)
                            print("   ✅ Gantt chart added to document")
                        else:
                            print("   ⚠️ Failed to generate base64 Gantt chart")
                    else:
                        print("   ⚠️ No Gantt data generated")
                except Exception as e:
                    print(f"   ❌ Error generating Gantt chart: {e}")
                    import traceback
                    traceback.print_exc()

        def _build_cover_section(doc, client_name: str):
            sec = doc.sections[0]
            sec.page_width = Cm(_PAGE_W_CM)
            sec.page_height = Cm(_PAGE_H_CM)
            sec.top_margin = Cm(1.0)
            sec.bottom_margin = Cm(1.0)
            sec.left_margin = Cm(1.75)
            sec.right_margin = Cm(1.75)

            bg = _asset("cover_bg")
            if bg:
                p_bg = doc.add_paragraph()
                p_bg.paragraph_format.space_after = Pt(0)
                _add_floating_picture(p_bg, bg, Cm(_PAGE_W_CM), Cm(_PAGE_H_CM), behind_doc=True, name="Cover Background")

            logo = _asset("logo_white")
            if logo:
                p_logo = doc.add_paragraph()
                p_logo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                p_logo.paragraph_format.space_after = Pt(0)
                r = p_logo.add_run()
                r.add_picture(logo, width=Cm(8.0))

            p_spacer = doc.add_paragraph()
            p_spacer.paragraph_format.space_after = Pt(210)

            p_rule = doc.add_paragraph()
            pPr = p_rule._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "4")
            bottom.set(qn("w:space"), "4")
            bottom.set(qn("w:color"), "FFFFFF")
            pBdr.append(bottom)
            pPr.append(pBdr)
            p_rule.paragraph_format.space_after = Pt(12)

            p_caption = doc.add_paragraph()
            p_caption.paragraph_format.space_after = Pt(14)
            _run(p_caption, "STATEMENT OF WORK BETWEEN", 15, bold=True, color=_PINK)

            p_company = doc.add_paragraph()
            p_company.paragraph_format.space_after = Pt(6)
            _run(p_company, client_name, 30, bold=True, color=_NAVY)

            p_and = doc.add_paragraph()
            p_and.paragraph_format.space_after = Pt(6)
            _run(p_and, "And", 16, italic=True, color=_GRAY)

            p_jman = doc.add_paragraph()
            p_jman.paragraph_format.space_after = Pt(0)
            _run(p_jman, "JMAN Group", 30, bold=True, color=_NAVY)

        def _build_content_section(doc, client_name="Client", version="v1.0"):
            sec = doc.add_section(WD_SECTION.NEW_PAGE)
            sec.page_width = Cm(_PAGE_W_CM)
            sec.page_height = Cm(_PAGE_H_CM)
            sec.top_margin = Cm(1.8)
            sec.bottom_margin = Cm(0.9)
            sec.left_margin = Cm(1.75)
            sec.right_margin = Cm(1.75)
            sec.header_distance = Cm(1.0)
            sec.footer_distance = Cm(0.9)
            sec.header.is_linked_to_previous = False
            sec.footer.is_linked_to_previous = False

            header = sec.header
            for p in list(header.paragraphs):
                p.clear()
            p_logo = header.paragraphs[0]
            p_logo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_logo.paragraph_format.space_after = Pt(0)
            logo = _asset("logo_navy")
            if logo:
                r = p_logo.add_run()
                r.add_picture(logo, width=Cm(3.2))

            footer = sec.footer
            for p in list(footer.paragraphs):
                p.clear()

            fp = footer.paragraphs[0]
            fp.paragraph_format.space_before = Pt(4)
            fp.paragraph_format.space_after = Pt(0)

            pPr = fp._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            top_bdr = OxmlElement("w:top")
            top_bdr.set(qn("w:val"), "single")
            top_bdr.set(qn("w:sz"), "8")
            top_bdr.set(qn("w:space"), "4")
            top_bdr.set(qn("w:color"), "173845")
            pBdr.append(top_bdr)
            pPr.append(pBdr)

            fp.paragraph_format.tab_stops.clear_all()
            fp.paragraph_format.tab_stops.add_tab_stop(Cm(0), WD_TAB_ALIGNMENT.LEFT)
            fp.paragraph_format.tab_stops.add_tab_stop(Cm(8.5), WD_TAB_ALIGNMENT.CENTER)
            fp.paragraph_format.tab_stops.add_tab_stop(Cm(17.5), WD_TAB_ALIGNMENT.RIGHT)

            month_year = datetime.now().strftime("%B %Y")
            _run(fp, f"Circulation Limited: {client_name}", 8, color=_GRAY)
            _run(fp, "\t", 8)
            _run(fp, f"Version {version}", 8, color=_GRAY)
            _run(fp, "\t", 8)
            _run(fp, month_year, 8, italic=True, color=_GRAY)
            _run(fp, " ", 8)
            _add_page_number_field(fp, 8, _GRAY)
            _run(fp, " ", 8)

            mark = _asset("corner_mark")
            if mark:
                r = fp.add_run()
                r.add_picture(mark, height=Cm(0.45))

            return sec

        # ── Main logic ────────────────────────────────────────────────────────

        messages = (
            db.query(Message)
            .filter(Message.conversation_id == conversation_id)
            .order_by(asc(Message.created_at))
            .all()
        )

        if not messages:
            raise HTTPException(status_code=404, detail="No messages found for this conversation.")

        # Find the last assistant message (the most recent generated proposal)
        last_assistant_message = None
        for msg in reversed(messages):
            if msg.role == "assistant":
                last_assistant_message = msg
                break

        if not last_assistant_message:
            raise HTTPException(status_code=404, detail="No assistant message found with proposal content.")

        last_proposal_content = last_assistant_message.content

        # Best-effort extract client name from first message
        client_name = "Client"
        try:
            first_message = messages[0] if messages else None
            if first_message and first_message.content:
                q_data = json.loads(first_message.content)
                client_name = (
                    q_data.get("company_name") or
                    q_data.get("client_name") or
                    q_data.get("organization") or
                    q_data.get("company") or
                    "Client"
                )
        except Exception:
            pass

        doc = Document()
        _build_cover_section(doc, client_name)
        _build_content_section(doc, client_name, "v1.0")

        heading_counter = [0]
        
        # Split the content into sections based on headings
        # We need to identify the Approach section by looking for "# Approach" heading
        sections = re.split(r'(?=^#\s+[A-Za-z])', last_proposal_content, flags=re.MULTILINE)
        
        for section in sections:
            section = section.strip()
            if not section:
                continue
            
            # Check if this is the Approach section
            is_approach = bool(re.match(r'^#\s+Approach\b', section, re.IGNORECASE))
            
            # Apply styled content with Gantt extraction for Approach section
            _apply_styled_content(doc, section, heading_counter, extract_gantt=is_approach)

        doc_stream = io.BytesIO()
        doc.save(doc_stream)
        doc_stream.seek(0)

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        clean_name = "".join(c for c in client_name if c.isalnum() or c in " ._-").strip()
        clean_name = clean_name.replace(" ", "_")
        filename = f"Proposal_{clean_name}_{timestamp}.docx"

        return StreamingResponse(
            doc_stream,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )

    except HTTPException:
        raise
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=400, detail=f"Error generating proposal document: {str(e)}")


def proposal_docx_upload(proposal_content: str, user_email: str, conversation_id: UUID):
    """Upload a generated proposal DOCX to blob storage and trigger search indexer"""
    try:
        doc = Document()
        title = doc.add_heading('AI Generated Proposal', 0)
        title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        add_content_to_doc(doc, proposal_content)

        doc_stream = io.BytesIO()
        doc.save(doc_stream)
        doc_stream.seek(0)

        container_name = GENERATED_PROPOSALS_CONTAINER_NAME or "workspaces"
        generated_blob_service_client = BlobServiceClient.from_connection_string(
            os.environ.get("AZURE_BLOB_STORAGE_CONNECTION_STRING")
        )
        blob_path = f"ai-proposals/{user_email}/{conversation_id}/generated_proposal_{uuid.uuid4()}.docx"
        generated_blob_client = generated_blob_service_client.get_blob_client(
            container=container_name,
            blob=blob_path
        )
        generated_blob_client.upload_blob(doc_stream, overwrite=True)
    except Exception as e:
        print(f"Warning: Could not upload proposal to blob storage: {e}")

    try:
        indexer_url = f'https://{AZURE_SEARCH_SERVICE}.search.windows.net/indexers/{AZURE_SEARCH_SERVICE_INDEXER_GENERATED_PROPOSALS_NAME}/run?api-version=2020-06-30'
        hdrs = {'api-key': AZURE_SEARCH_KEY, 'Content-Type': 'application/json'}
        response = requests.post(indexer_url, headers=hdrs)
        response.raise_for_status()

        indexer_status_url = f'https://{AZURE_SEARCH_SERVICE}.search.windows.net/indexers/{AZURE_SEARCH_SERVICE_INDEXER_GENERATED_PROPOSALS_NAME}/status?api-version=2020-06-30'
        max_retries = 10
        wait_time = 5
        last_run_status = "Unknown"

        for _ in range(max_retries):
            status_response = requests.get(indexer_status_url, headers=hdrs)
            status_data = status_response.json()
            last_run_status = status_data.get("lastResult", {}).get("status", "Unknown")
            if last_run_status in ["success", "transientFailure", "permanentFailure"]:
                break
            time.sleep(wait_time)
    except Exception:
        pass


async def upload_files(conversation_id: UUID, files: List[UploadFile], db: Session, user: User):
    """Upload multiple files to blob storage and trigger search indexer"""
    uploaded_metadata = []
    last_run_status = "Unknown"
    last_run_error = ""

    for file in files:
        contents = await file.read()
        blob_path = f"{user.email}----{conversation_id}/{user.email}----{conversation_id}----{file.filename}"
        await upload_file_to_azure_blob(contents, blob_path)
        uploaded_metadata.append({"filename": file.filename})

    try:
        indexer_url = f'https://{AZURE_SEARCH_SERVICE}.search.windows.net/indexers/{AZURE_SEARCH_AI_PROPOSAL_WORKSPACE_INDEXER}/run?api-version=2020-06-30'
        hdrs = {'api-key': AZURE_SEARCH_KEY, 'Content-Type': 'application/json'}
        response = requests.post(indexer_url, headers=hdrs)
        response.raise_for_status()

        indexer_status_url = f'https://{AZURE_SEARCH_SERVICE}.search.windows.net/indexers/{AZURE_SEARCH_AI_PROPOSAL_WORKSPACE_INDEXER}/status?api-version=2020-06-30'
        max_retries = 10
        wait_time = 5

        for _ in range(max_retries):
            status_response = requests.get(indexer_status_url, headers=hdrs)
            status_data = status_response.json()
            last_run_status = status_data.get("lastResult", {}).get("status", "Unknown")
            last_run_error = status_data.get("lastResult", {}).get("errorMessage", "")
            if last_run_status in ["success", "transientFailure", "permanentFailure"]:
                break
            time.sleep(wait_time)
    except Exception:
        pass
    finally:
        return {
            "success": True,
            "uploaded": uploaded_metadata,
            "indexer_status": last_run_status,
            "error_message": last_run_error,
            "status": last_run_status == "success"
        }


async def upload_file_to_azure_blob(file: bytes, blob_path: str):
    """Upload a single file to Azure Blob Storage"""
    try:
        blob_service_client_async = AsyncBlobServiceClient.from_connection_string(
            os.environ.get("AZURE_BLOB_STORAGE_CONNECTION_STRING")
        )
        container_name = AI_PROPOSAL_WORKSPACE_CONTAINER_NAME or "workspaces"
        container_client = blob_service_client_async.get_container_client(container_name)
        blob_client = container_client.get_blob_client(f"ai-proposals/{blob_path}")
        await blob_client.upload_blob(file, overwrite=True)
    except Exception as e:
        print(f"Warning: Could not upload file to blob storage: {e}")
        raise HTTPException(status_code=400, detail=f"File upload failed: {str(e)}")


async def edit_proposal_llm_service(conversation_id: UUID, new_message: str, db: Session, current_user: User, message_id: UUID = None):
    """
    Edit a specific part of a proposal using AI based on user instructions.
    If message_id is provided, uses that specific message; otherwise uses the latest proposal.
    """
    if message_id:
        message = db.query(Message).filter(
            Message.conversation_id == conversation_id,
            Message.id == message_id
        ).first()
        if not message:
            raise HTTPException(status_code=404, detail="Message not found")
        proposal_msg = message.content
    else:
        # Get latest proposal
        messages = (
            db.query(Message)
            .filter(Message.conversation_id == conversation_id, Message.role == "assistant")
            .order_by(Message.created_at.desc())
            .first()
        )
        if not messages:
            raise HTTPException(status_code=404, detail="No proposal found")
        proposal_msg = messages.content
    
    system_prompt = f"""
User will provide a part of proposal's content along with new instructions to edit it.
You need to update ONLY the part of the proposal that is specified in the new instructions by taking context of whole proposal.

Whole proposal:
{proposal_msg}
"""
    messages = [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": new_message}
    ]

    try:
        response = azure_client.chat.completions.create(
            model=AZURE_OPENAI_DEPLOYMENT,
            messages=messages,
            temperature=float(AZURE_OPENAI_TEMPERATURE),
            max_tokens=10000,
            top_p=float(AZURE_OPENAI_TOP_P),
        )
        
        return {
            "edited_proposal": response.choices[0].message.content
        }
        
    except Exception as e:
        print(f"Azure OpenAI error: {e}")
        raise HTTPException(status_code=500, detail=f"AI generation failed: {str(e)}")